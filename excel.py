import logging

from openpyxl import load_workbook

from models import WorkShift


logger = logging.getLogger(__name__)


def write_to_excel(
    shifts: list[WorkShift],
    template_path: str,
    output_path: str,
    start_row: int = 9,
) -> None:
    """
    Write parsed shifts to an Excel workbook.

    Data is written to every second column:
    A, C, E, G, I, K...
    """

    try:
        workbook = load_workbook(template_path)

    except FileNotFoundError:
        logger.error("Excel template not found: %s", template_path)
        raise

    sheet = workbook.active

    for row_index, shift in enumerate(shifts, start=start_row):

        values = [
            shift.date,
            shift.start,
            shift.end,
            shift.city,
            shift.hours,
            shift.prevailing_hours,
        ]

        for col_index, value in enumerate(values, start=1):
            column = col_index * 2 - 1

            sheet.cell(
                row=row_index,
                column=column,
                value=value,
            )

    try:
        workbook.save(output_path)

    except PermissionError:
        logger.error(
            "Could not save Excel file. " "Make sure the file is not open: %s",
            output_path,
        )
        raise

    logger.info("Excel file saved: %s", output_path)
