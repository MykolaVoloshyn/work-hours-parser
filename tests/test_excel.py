from openpyxl import Workbook, load_workbook

from excel import write_to_excel
from models import WorkShift


def test_write_to_excel(tmp_path):
    # Create temporary Excel template
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "result.xlsx"

    workbook = Workbook()
    sheet = workbook.active

    workbook.save(template_path)

    # Test data
    shifts = [
        WorkShift(
            date="1/03",
            start="8:00",
            end="3:00",
            city="Kirkland",
            hours=7.0,
        ),
        WorkShift(
            date="3/03",
            start="6:00",
            end="5:30",
            city="Whidbey, Ferndale",
            hours=11.5,
            prevailing_hours=3.5,
        ),
    ]

    # Run function
    write_to_excel(
        shifts=shifts,
        template_path=str(template_path),
        output_path=str(output_path),
    )

    # Open generated file
    result_workbook = load_workbook(output_path)
    result_sheet = result_workbook.active

    # First row
    assert result_sheet["A9"].value == "1/03"
    assert result_sheet["C9"].value == "8:00"
    assert result_sheet["E9"].value == "3:00"
    assert result_sheet["G9"].value == "Kirkland"
    assert result_sheet["I9"].value == 7.0
    assert result_sheet["K9"].value is None

    # Second row
    assert result_sheet["A10"].value == "3/03"
    assert result_sheet["C10"].value == "6:00"
    assert result_sheet["E10"].value == "5:30"
    assert result_sheet["G10"].value == "Whidbey, Ferndale"
    assert result_sheet["I10"].value == 11.5
    assert result_sheet["K10"].value == 3.5
